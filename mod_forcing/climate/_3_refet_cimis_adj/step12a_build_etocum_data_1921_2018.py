"""Build a standalone 1921-2018 EToCum workbook for the spline workflow."""

import pathlib
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

# Add repo root to path for utils imports
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[3]))
from utils.paths import get_module_generated_dir


HERE = pathlib.Path(__file__).resolve().parent
gen_dir = get_module_generated_dir("mod_forcing/climate")
STEP11_DIR = gen_dir / "output" / "_3_refet_cimis_adj" / "step11"
OUTPUT_DIR = gen_dir / "output" / "_3_refet_cimis_adj" / "step12"

TIMESERIES_FILES = {
    "WBA": STEP11_DIR / "Historical_Unsplit (Product_A)_Spatial_Average_ETo_F1corrected_monthly_timeseries_per_WBA.csv",
    "SR": STEP11_DIR / "Historical_Unsplit (Product_A)_Spatial_Average_ETo_F1corrected_monthly_timeseries_per_SR.csv",
    "SW": STEP11_DIR / "Historical_Unsplit (Product_A)_Spatial_Average_ETo_F1corrected_monthly_timeseries_per_SW.csv",
}
TEMPLATE_XLSM = HERE.parent / "reference" / "CS3_Table_9-xx_ETo_Spline2025_Oct1920.xlsm"
OUTPUT_XLSX = OUTPUT_DIR / "Historical_Unsplit (Product_A)_EToCum_1921_2018.xlsx"

SHEET_NAME = "EToCum"
HEADER_ROW = 2
DATE_COLUMN = 3
DATA_START_COLUMN = 4
FIRST_DATA_ROW = 8
START_YEAR = 1921
END_YEAR = 2018
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

def expected_month_columns():
    return [f"{month}-{year}" for year in range(START_YEAR, END_YEAR + 1)
            for month in MONTHS]


def expected_dates():
    return pd.date_range(
        f"{START_YEAR}-01-01",
        periods=(END_YEAR - START_YEAR + 1) * 12,
        freq=pd.offsets.MonthEnd(),
    )


def load_spatial_timeseries():
    month_columns = expected_month_columns()
    id_columns = {"WBA": "WBA_ID", "SR": "SR_ID", "SW": "SW_ID"}
    tables = {}
    for group, path in TIMESERIES_FILES.items():
        id_column = id_columns[group]
        if not path.is_file():
            raise FileNotFoundError(f"Step 11 {group} monthly_timeseries file not found:\n  {path}")
        df = pd.read_csv(path, dtype={id_column: str})
        expected_columns = [id_column] + month_columns
        if list(df.columns) != expected_columns:
            raise ValueError(
                f"{path.name} must contain {id_column} followed by the exact "
                f"{START_YEAR}-{END_YEAR} monthly columns."
            )
        if df[id_column].duplicated().any():
            raise ValueError(f"{path.name} contains duplicate {id_column} values.")
        values = df[month_columns].to_numpy(dtype=float)
        if not np.isfinite(values).all():
            raise ValueError(f"{path.name} contains missing or nonnumeric monthly values.")
        tables[group] = df.set_index(id_column)[month_columns]
    return tables, month_columns


def source_id_for_header(header):
    if header.startswith("WBA_SR_"):
        return "SR", "SR_" + header.removeprefix("WBA_SR_")
    if header.startswith("WBA_SW_"):
        return "SW", "SW_" + header.removeprefix("WBA_SW_")
    if header.startswith("WBA_"):
        return "WBA", header.removeprefix("WBA_")
    raise ValueError(f"Unsupported EToCum series header: {header!r}")


def read_template_layout():
    if not TEMPLATE_XLSM.is_file():
        raise FileNotFoundError(f"Spline workbook template not found:\n  {TEMPLATE_XLSM}")

    workbook = load_workbook(TEMPLATE_XLSM, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise KeyError(f"Workbook has no {SHEET_NAME!r} sheet.")
        sheet = workbook[SHEET_NAME]
        header_values = next(sheet.iter_rows(
            min_row=HEADER_ROW, max_row=HEADER_ROW,
            min_col=DATA_START_COLUMN, max_col=sheet.max_column,
            values_only=True,
        ))
        header_columns = {
            str(value): column
            for column, value in enumerate(header_values, start=DATA_START_COLUMN)
            if value is not None and str(value).strip()
        }
        header_count = sum(
            1 for value in header_values
            if value is not None and str(value).strip()
        )
        if len(header_columns) != header_count:
            raise ValueError("EToCum row 2 contains duplicate series headers.")

        target_date_rows = {}
        date_values = sheet.iter_rows(
            min_row=FIRST_DATA_ROW, max_row=sheet.max_row,
            min_col=DATE_COLUMN, max_col=DATE_COLUMN,
            values_only=True,
        )
        for row_number, (value,) in enumerate(date_values, start=FIRST_DATA_ROW):
            if value is None:
                continue
            date = pd.Timestamp(value).normalize()
            if START_YEAR <= date.year <= END_YEAR:
                if date in target_date_rows:
                    raise ValueError(f"Duplicate EToCum date row: {date.date()}")
                target_date_rows[date] = row_number
    finally:
        workbook.close()

    dates = expected_dates()
    if list(target_date_rows) != list(dates):
        missing = [date for date in dates if date not in target_date_rows]
        extra = [date for date in target_date_rows if date not in set(dates)]
        raise ValueError(
            f"EToCum must contain every month-end from {START_YEAR} through {END_YEAR}; "
            f"missing={missing[:3]}, extra={extra[:3]}"
        )
    if "WBA_91" not in header_columns:
        raise ValueError("EToCum is missing the required WBA_91 column.")
    return header_columns, target_date_rows, dates


def load_weights(path, header):
    if not path.is_file():
        raise FileNotFoundError(f"Step 0 weight table for {header} not found:\n  {path}")
    weights = pd.read_csv(path, skiprows=2).rename(columns={
        "long[k]": "Lon",
        "lati[k]": "Lat",
        "percent[k]": "Weight",
    })
    required = {"Lon", "Lat", "Weight"}
    if not required.issubset(weights.columns):
        raise ValueError(f"Invalid Step 0 weight table columns for {header}: {path.name}")
    weights = weights[["Lon", "Lat", "Weight"]].copy()
    weights[["Lon", "Lat"]] = weights[["Lon", "Lat"]].round(7)
    weights["Weight"] = pd.to_numeric(weights["Weight"], errors="coerce")
    if weights.duplicated(["Lon", "Lat"]).any():
        raise ValueError(f"{header}: duplicate grid coordinates in {path.name}")
    values = weights["Weight"].to_numpy(dtype=float)
    if not np.isfinite(values).all() or (values < 0).any() or values.sum() <= 0:
        raise ValueError(f"{header}: weights must be finite, nonnegative, and nonzero.")
    if values.sum() < 95:
        print(f"[WARN] {header}: Step 0 weights cover {values.sum():.2f}% of polygon area.")
    return weights


def weighted_monthly_series(
    source_values, source_index, month_columns, weight_path, header,
):
    weights = load_weights(weight_path, header)
    weight_index = pd.MultiIndex.from_frame(weights[["Lon", "Lat"]])
    positions = source_index.get_indexer(weight_index)
    unmatched = positions < 0
    if unmatched.any():
        examples = weights.loc[unmatched, ["Lon", "Lat"]].head(3).to_dict("records")
        raise ValueError(
            f"{header}: {int(unmatched.sum())} Step 0 coordinates are absent from Step 10b; "
            f"examples: {examples}"
        )

    weight_values = weights["Weight"].to_numpy(dtype=float)
    values = source_values[positions, :]
    valid = np.isfinite(values)
    weighted_totals = np.where(valid, values, 0.0).T @ weight_values
    valid_weights = valid.T @ weight_values
    series = np.divide(
        weighted_totals,
        valid_weights,
        out=np.full(len(month_columns), np.nan),
        where=valid_weights > 0,
    )
    if not np.isfinite(series).all():
        missing_months = [month_columns[i] for i in np.flatnonzero(~np.isfinite(series))[:5]]
        raise ValueError(f"{header}: missing weighted values for months {missing_months}")
    return series


def write_etocum_workbook(result):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(
        OUTPUT_XLSX,
        engine="openpyxl",
        date_format="mmm yyyy",
        datetime_format="mmm yyyy",
    ) as writer:
        result.to_excel(writer, sheet_name=SHEET_NAME, index=False)
        sheet = writer.sheets[SHEET_NAME]
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 14
        for row in range(2, len(result) + 2):
            sheet.cell(row=row, column=1).number_format = "m/d/yyyy"


def main():
    tables, _ = load_spatial_timeseries()
    header_columns, target_date_rows, dates = read_template_layout()
    series_by_header = {}

    for header in header_columns:
        if header == "WBA_91":
            series_by_header[header] = None
            continue
        group, source_id = source_id_for_header(header)
        if source_id not in tables[group].index:
            raise ValueError(f"{header}: no matching Step 11 {group} series {source_id!r}.")
        series_by_header[header] = tables[group].loc[source_id].to_numpy(dtype=float)

    blank_headers = [header for header, values in series_by_header.items() if values is None]
    if blank_headers != ["WBA_91"]:
        raise ValueError(f"WBA_91 must be the only blank series; got {blank_headers}.")

    result_data = {"Date": dates}
    for header in header_columns:
        series = series_by_header[header]
        result_data[header] = np.nan if series is None else series
    result = pd.DataFrame(result_data)

    write_etocum_workbook(result)
    workbook = load_workbook(OUTPUT_XLSX, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [SHEET_NAME]:
            raise ValueError(f"Output must contain only the {SHEET_NAME!r} sheet.")
        sheet = workbook[SHEET_NAME]
        actual_headers = next(sheet.iter_rows(
            min_row=1, max_row=1, values_only=True,
        ))
        expected_headers = ["Date"] + list(header_columns)
        if list(actual_headers) != expected_headers:
            raise ValueError("Output EToCum workbook headers do not match the template order.")
        actual_rows = list(sheet.iter_rows(
            min_row=2, max_row=len(dates) + 1,
            max_col=len(expected_headers), values_only=True,
        ))
        if len(actual_rows) != len(dates):
            raise ValueError(f"Output has {len(actual_rows)} rows; expected {len(dates)}.")
        if any(row[1 + list(header_columns).index("WBA_91")] is not None for row in actual_rows):
            raise ValueError("WBA_91 must remain blank in the output workbook.")
        for row_index, row in enumerate(actual_rows):
            if pd.Timestamp(row[0]).normalize() != dates[row_index]:
                raise ValueError(f"Output date mismatch on row {row_index + 2}.")
            if any(value is None for column, value in enumerate(row[1:], start=1)
                   if expected_headers[column] != "WBA_91"):
                raise ValueError(f"Output has a blank required value on row {row_index + 2}.")
    finally:
        workbook.close()

    print(f"EToCum-only workbook saved: {OUTPUT_XLSX}")
    print(f"Rows: {len(dates)} ({dates[0].date()} through {dates[-1].date()})")
    print(f"Series populated: {len(header_columns) - 1}; WBA_91 left blank.")
    print("This output contains only EToCum data; it is not the spline macro workbook.")


if __name__ == "__main__":
    main()